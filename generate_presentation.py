import sys
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 widescreen
    prs.slide_height = Inches(7.5)

    # Color Constants
    PURPLE_PRIMARY = RGBColor(109, 40, 217)   # #6D28D9
    PURPLE_DARK = RGBColor(76, 29, 149)      # #4C1D95
    PURPLE_LIGHT = RGBColor(245, 243, 255)   # #F5F3FF
    PURPLE_BORDER = RGBColor(221, 214, 254)  # #DDD6FE
    SLATE_DARK = RGBColor(15, 23, 42)        # #0F172A
    SLATE_MUTED = RGBColor(71, 85, 105)      # #475569
    WHITE = RGBColor(255, 255, 255)
    EMERALD = RGBColor(16, 185, 129)
    AMBER = RGBColor(245, 158, 11)
    ROSE = RGBColor(239, 68, 68)

    blank_layout = prs.slide_layouts[6]

    def add_header(slide, category, title):
        # Top banner background shape
        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
        banner.fill.solid()
        banner.fill.fore_color.rgb = PURPLE_DARK
        banner.line.color.rgb = PURPLE_PRIMARY

        # Category text
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.7), Inches(0.3))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = category.upper()
        p_c.font.name = 'Arial'
        p_c.font.size = Pt(10)
        p_c.font.bold = True
        p_c.font.color.rgb = RGBColor(216, 180, 254)

        # Title text
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.65))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.name = 'Arial'
        p_t.font.size = Pt(20)
        p_t.font.bold = True
        p_t.font.color.rgb = WHITE

    def add_card(slide, left, top, width, height, bg_color=PURPLE_LIGHT, border_color=PURPLE_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
        return card

    # ====================================================
    # SLIDE 1: TITLE / COVER SLIDE
    # ====================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = PURPLE_DARK
    bg1.line.fill.background()

    c1 = add_card(s1, 1.2, 0.9, 10.933, 5.7, WHITE, PURPLE_PRIMARY)

    tb1 = s1.shapes.add_textbox(Inches(1.8), Inches(1.3), Inches(9.733), Inches(4.9))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p1_sub = tf1.paragraphs[0]
    p1_sub.text = "MEDIAI TELEHEALTH PLATFORM • EXECUTIVE SPECIFICATION"
    p1_sub.font.name = 'Arial'
    p1_sub.font.size = Pt(11)
    p1_sub.font.bold = True
    p1_sub.font.color.rgb = PURPLE_PRIMARY

    p1_title = tf1.add_paragraph()
    p1_title.text = "Pre-Consultation Safety Triage & Encrypted Telehealth Engine"
    p1_title.font.name = 'Arial'
    p1_title.font.size = Pt(28)
    p1_title.font.bold = True
    p1_title.font.color.rgb = PURPLE_DARK
    p1_title.space_before = Pt(6)
    p1_title.space_after = Pt(10)

    p1_desc = tf1.add_paragraph()
    p1_desc.text = (
        "A Unified Healthcare Ecosystem Combining Multi-Tier Differential Diagnosis, "
        "Patient Safety Contraindication Protocols, WhatsApp-Grade 256-Bit Encrypted Messaging, "
        "Verified Physician Video Demonstration Hub, and Medical Board Credential Verification."
    )
    p1_desc.font.name = 'Arial'
    p1_desc.font.size = Pt(12.5)
    p1_desc.font.color.rgb = SLATE_DARK
    p1_desc.space_after = Pt(18)

    p1_meta = tf1.add_paragraph()
    p1_meta.text = (
        "Lead Engineer: Daniel Isaac Echo (CS/DE/2022/1300)\n"
        "Department of Computer Science • Caritas University, Enugu\n"
        "Supervisor & HOD: Mr. Ejike Ugwu • Project Coordinator: Mr. Abundance Akpan\n"
        "Date: September 2026 • Version 2.5 (Clinical Safety Edition)"
    )
    p1_meta.font.name = 'Arial'
    p1_meta.font.size = Pt(10)
    p1_meta.font.color.rgb = SLATE_MUTED

    # ====================================================
    # SLIDE 2: THE CLINICAL PROBLEM & THE PHARMACY MISCONCEPTION
    # ====================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Context & Clinical Need", "The Problem: Hazardous Self-Medication & Autonomous Prescribing Risks")

    prob_cards = [
        ("The Autonomous Pharmacist Trap", "A common mistake in AI healthcare is attempting to build an autonomous pharmacist that dispenses drugs directly to patients. This creates extreme medical malpractice hazards, because algorithms cannot perform physical auscultation, palpation, or endoscopy.", 0.8, 1.5, 3.733, 5.3),
        ("The Lethal NSAID Contraindication", "In real-world clinics, patients with gastric ulcers or severe stomach irritation often feel pain referring to their back. If an AI presumes body pain and prescribes NSAIDs (Ibuprofen, Diclofenac, Aspirin), it triggers severe mucosal erosion and life-threatening internal bleeding!", 4.8, 1.5, 3.733, 5.3),
        ("The MediAI Paradigm Shift", "Instead of autonomous drug prescribing, MediAI delivers Pre-Consultation Safety Triage: guiding patients on what to do and NOT to do (strictly avoiding NSAIDs, taking plain Paracetamol, chair transition posture) and bridging them seamlessly into consultations with verified doctors.", 8.8, 1.5, 3.733, 5.3),
    ]

    for title, text, l, t, w, h in prob_cards:
        add_card(s2, l, t, w, h, WHITE, PURPLE_PRIMARY)
        tb = s2.shapes.add_textbox(Inches(l + 0.25), Inches(t + 0.25), Inches(w - 0.5), Inches(h - 0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p_h = tf.paragraphs[0]
        p_h.text = title
        p_h.font.name = 'Arial'
        p_h.font.size = Pt(13)
        p_h.font.bold = True
        p_h.font.color.rgb = PURPLE_DARK
        p_h.space_after = Pt(10)

        for line in text.split('\n'):
            p_b = tf.add_paragraph()
            p_b.text = line
            p_b.font.name = 'Arial'
            p_b.font.size = Pt(10)
            p_b.font.color.rgb = SLATE_DARK
            p_b.space_after = Pt(4)

    # ====================================================
    # SLIDE 3: CLINICAL CASE STUDY (THE REAL-LIFE CASE)
    # ====================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Clinical Case Study", "Real-Life Validation: The Ulcer-Spine Referral Diagnostic Architecture")

    case_cards = [
        ("The Presenting Symptom", "Patient 'Baby ❤️' presented with severe mid-back pain that felt internally like a previous ulcer episode. Key mechanical symptom: whenever she sat for a while and attempted to stand, her back locked up in severe pain, preventing her from standing straight.", 0.8, 1.5, 3.733, 5.3),
        ("Differential Mechanisms", "• Referred Gastric Ulcer Pain: Posterior stomach wall lesions refer pain directly to mid-back nerves (T5-T10). Sitting compresses the stomach; standing stretches the wall, triggering locking spasms.\n• Psoas Muscle Spasm: Shortens when sitting; forced stretch on standing triggers sharp lock.\n• Biliary / Pancreatic Referral: Deep epigastric ache.", 4.8, 1.5, 3.733, 5.3),
        ("Pre-Consultation Safety Directives", "• STRICTLY AVOID NSAIDs: Zero Ibuprofen, Diclofenac, Feldene, Aspirin (avoids gastric hemorrhage).\n• SAFE OTC OPTION: Plain Paracetamol (Acetaminophen) as the sole safe pain reliever.\n• ERGONOMIC CARE: Sip warm water, avoid peppery food, brace on chair edge and stand gradually.", 8.8, 1.5, 3.733, 5.3),
    ]

    for title, text, l, t, w, h in case_cards:
        add_card(s3, l, t, w, h, PURPLE_LIGHT, PURPLE_BORDER)
        tb = s3.shapes.add_textbox(Inches(l + 0.25), Inches(t + 0.25), Inches(w - 0.5), Inches(h - 0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p_h = tf.paragraphs[0]
        p_h.text = title
        p_h.font.name = 'Arial'
        p_h.font.size = Pt(13)
        p_h.font.bold = True
        p_h.font.color.rgb = PURPLE_PRIMARY
        p_h.space_after = Pt(8)

        for line in text.split('\n'):
            p_b = tf.add_paragraph()
            p_b.text = line
            p_b.font.name = 'Arial'
            p_b.font.size = Pt(9.5)
            p_b.font.color.rgb = SLATE_DARK
            p_b.space_after = Pt(4)

    # ====================================================
    # SLIDE 4: SYSTEM ARCHITECTURE
    # ====================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Engineering Specifications", "Multi-Tier Software Architecture & Technology Stack")

    arch_tiers = [
        ("Presentation Layer", "React 18 SPA + Vite\nTailwind CSS Custom Design\nZustand State Management\nClean Markdown Parser (No # Tags)", 0.8, 1.5, 2.7, 5.3),
        ("API & Gateway", "Django 6.1 + DRF\nDaphne ASGI Server\nJWT Bearer Authentication\nRole-Based Access Control", 3.8, 1.5, 2.7, 5.3),
        ("Asynchronous Engine", "Django Channels 4.3\nBidirectional WebSockets\nEncrypted WhatsApp Messenger\nLive Triage Broadcasting", 6.8, 1.5, 2.7, 5.3),
        ("Persistence & Media", "Relational Database Schema\nJSONB Clinical Payloads\nEncrypted File Storage\nDirect MP4/WebM Video Stream", 9.8, 1.5, 2.7, 5.3),
    ]

    for title, text, l, t, w, h in arch_tiers:
        add_card(s4, l, t, w, h, WHITE, PURPLE_PRIMARY)
        tb = s4.shapes.add_textbox(Inches(l + 0.2), Inches(t + 0.25), Inches(w - 0.4), Inches(h - 0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p_h = tf.paragraphs[0]
        p_h.text = title
        p_h.font.name = 'Arial'
        p_h.font.size = Pt(13)
        p_h.font.bold = True
        p_h.font.color.rgb = PURPLE_DARK
        p_h.space_after = Pt(12)

        for line in text.split('\n'):
            p_b = tf.add_paragraph()
            p_b.text = "• " + line
            p_b.font.name = 'Arial'
            p_b.font.size = Pt(10)
            p_b.font.color.rgb = SLATE_DARK
            p_b.space_after = Pt(6)

    # ====================================================
    # SLIDE 5: WHATSAPP-GRADE SECURE CLINICAL MESSENGER
    # ====================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Clinical Communication", "WhatsApp-Grade 256-Bit Encrypted Clinical Consultation Messenger")

    chat_cards = [
        ("256-Bit End-to-End Encryption Banner", "Ensures compliance with HIPAA and GDPR medical privacy standards. Patients and clinicians communicate over authenticated channels with complete confidentiality.", 0.8, 1.5, 3.733, 5.3),
        ("Double Blue Checkmarks (Read Receipts)", "Real-time message status indicators (CheckCheck icon) confirm instant transmission, delivery, and review by the attending medical specialist.", 4.8, 1.5, 3.733, 5.3),
        ("1-Click Clinical Triage Sharing & Video Launch", "Patients can transmit their AI symptom assessment directly into the chat thread with one click. The chat header includes a direct 'Video Call' button to launch the WebRTC teleconsultation room instantly.", 8.8, 1.5, 3.733, 5.3),
    ]

    for title, text, l, t, w, h in chat_cards:
        add_card(s5, l, t, w, h, WHITE, PURPLE_PRIMARY)
        tb = s5.shapes.add_textbox(Inches(l + 0.25), Inches(t + 0.25), Inches(w - 0.5), Inches(h - 0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p_h = tf.paragraphs[0]
        p_h.text = title
        p_h.font.name = 'Arial'
        p_h.font.size = Pt(13)
        p_h.font.bold = True
        p_h.font.color.rgb = PURPLE_DARK
        p_h.space_after = Pt(10)

        for line in text.split('\n'):
            p_b = tf.add_paragraph()
            p_b.text = line
            p_b.font.name = 'Arial'
            p_b.font.size = Pt(10)
            p_b.font.color.rgb = SLATE_DARK
            p_b.space_after = Pt(4)

    # ====================================================
    # SLIDE 6: FLO-INSPIRED HEALTH COMMUNITY & VIDEO DEMONSTRATIONS
    # ====================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Social Engagement & Education", "Flo-Inspired Verified Medical Creator & Video Demonstration Hub")

    comm_cards = [
        ("Doctor Video Demonstrations", "Doctors can upload MP4/WebM video demonstrations directly from their devices—showing correct blood pressure measurement, physical therapy exercises, and dietary protocols for patients to watch on demand.", 0.8, 1.5, 3.733, 5.3),
        ("Clean Typography (No Raw # Hashtags)", "The article reader transforms raw clinical markdown into clean, magazine-grade typography: styled purple section headers, structured bullet lists, and highlighted clinical pearl cards.", 4.8, 1.5, 3.733, 5.3),
        ("Strict Creator Gating & Monetization", "Only approved doctors and verified wellness specialists have permission to publish. The Creator Studio provides private analytics ($0.02/read view, $0.05/like) and direct bank payout configuration.", 8.8, 1.5, 3.733, 5.3),
    ]

    for title, text, l, t, w, h in comm_cards:
        add_card(s6, l, t, w, h, PURPLE_LIGHT, PURPLE_BORDER)
        tb = s6.shapes.add_textbox(Inches(l + 0.25), Inches(t + 0.25), Inches(w - 0.5), Inches(h - 0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p_h = tf.paragraphs[0]
        p_h.text = title
        p_h.font.name = 'Arial'
        p_h.font.size = Pt(13)
        p_h.font.bold = True
        p_h.font.color.rgb = PURPLE_PRIMARY
        p_h.space_after = Pt(8)

        for line in text.split('\n'):
            p_b = tf.add_paragraph()
            p_b.text = line
            p_b.font.name = 'Arial'
            p_b.font.size = Pt(9.5)
            p_b.font.color.rgb = SLATE_DARK
            p_b.space_after = Pt(4)

    # ====================================================
    # SLIDE 7: DOCTOR CREDENTIAL VERIFICATION & ADMIN AUDIT
    # ====================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Governance & Quality Control", "Medical Board Credential Verification & Administrative Audit Portal")

    gov_cards = [
        ("Doctor Verification Portal", "Doctors upload their official Medical Practicing Licence, MBBS/MD Certificate, and Government-Issued ID directly in their dashboard. Status remains 'In Review' until authorized.", 0.8, 1.5, 3.733, 5.3),
        ("Admin Credential Inspection Modal", "Platform administrators inspect submitted documents via a dedicated audit modal—verifying licence numbers, document types, and issuing dates before approving practice rights.", 4.8, 1.5, 3.733, 5.3),
        ("1-Click Approval & Revocation", "Administrators can approve verified doctors with a single click (granting consultation, video, and publishing permissions) or revoke access instantly if credentials lapse.", 8.8, 1.5, 3.733, 5.3),
    ]

    for title, text, l, t, w, h in gov_cards:
        add_card(s7, l, t, w, h, WHITE, PURPLE_PRIMARY)
        tb = s7.shapes.add_textbox(Inches(l + 0.25), Inches(t + 0.25), Inches(w - 0.5), Inches(h - 0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p_h = tf.paragraphs[0]
        p_h.text = title
        p_h.font.name = 'Arial'
        p_h.font.size = Pt(13)
        p_h.font.bold = True
        p_h.font.color.rgb = PURPLE_DARK
        p_h.space_after = Pt(10)

        for line in text.split('\n'):
            p_b = tf.add_paragraph()
            p_b.text = line
            p_b.font.name = 'Arial'
            p_b.font.size = Pt(10)
            p_b.font.color.rgb = SLATE_DARK
            p_b.space_after = Pt(4)

    # ====================================================
    # SLIDE 8: CONCLUSION & EXECUTIVE SUMMARY
    # ====================================================
    s8 = prs.slides.add_slide(blank_layout)
    bg8 = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg8.fill.solid()
    bg8.fill.fore_color.rgb = PURPLE_DARK
    bg8.line.fill.background()

    c8 = add_card(s8, 1.2, 1.0, 10.933, 5.5, WHITE, PURPLE_PRIMARY)

    tb8 = s8.shapes.add_textbox(Inches(1.8), Inches(1.4), Inches(9.733), Inches(4.7))
    tf8 = tb8.text_frame
    tf8.word_wrap = True

    p8_sub = tf8.paragraphs[0]
    p8_sub.text = "EXECUTIVE SUMMARY & DEMONSTRATION READINESS"
    p8_sub.font.name = 'Arial'
    p8_sub.font.size = Pt(12)
    p8_sub.font.bold = True
    p8_sub.font.color.rgb = PURPLE_PRIMARY

    p8_t = tf8.add_paragraph()
    p8_t.text = "MediAI: Transforming Telehealth Through Clinical Intelligence & Safety"
    p8_t.font.name = 'Arial'
    p8_t.font.size = Pt(26)
    p8_t.font.bold = True
    p8_t.font.color.rgb = PURPLE_DARK
    p8_t.space_before = Pt(6)
    p8_t.space_after = Pt(12)

    p8_body = tf8.add_paragraph()
    p8_body.text = (
        "1. Corrected Core Aim: Operates as an ethical Pre-Consultation Safety Triage & Doctor Bridge, eliminating dangerous autonomous drug dispensing.\n"
        "2. Clinical Case Validated: Proven differential analysis of real-world ulcer-spine referral cases with strict NSAID contraindication warnings.\n"
        "3. WhatsApp-Grade Messenger: 256-bit encrypted clinical chat with double blue checkmarks and 1-click triage memo sharing.\n"
        "4. Video Health Community: Direct physician video demonstration streaming and clean semantic articles without raw hashtag symbols.\n"
        "5. Medical Governance: Rigorous credential verification pipeline protecting patient safety."
    )
    p8_body.font.name = 'Arial'
    p8_body.font.size = Pt(11.5)
    p8_body.font.color.rgb = SLATE_DARK
    p8_body.space_after = Pt(14)

    prs.save("MediAI_Executive_Presentation.pptx")
    print("Saved: MediAI_Executive_Presentation.pptx")

if __name__ == "__main__":
    build_presentation()
