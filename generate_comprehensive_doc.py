import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_comprehensive_documentation():
    doc = Document()

    # Set page margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    PURPLE_PRIMARY = RGBColor(109, 40, 217)    # #6D28D9
    PURPLE_DARK = RGBColor(76, 29, 149)       # #4C1D95
    SLATE_DARK = RGBColor(15, 23, 42)         # #0F172A
    SLATE_MUTED = RGBColor(71, 85, 105)       # #475569
    HEX_PURPLE_LIGHT = "F5F3FF"
    HEX_HEADER_BG = "4C1D95"
    HEX_ALT_ROW = "F8FAFC"

    def set_cell_background(cell, hex_color):
        shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
        cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    # ----------------------------------------------------
    # COVER TITLE & HEADER
    # ----------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(20)
    title_p.paragraph_format.space_after = Pt(4)
    run_sub = title_p.add_run("MEDIAI TELEHEALTH PLATFORM • SYSTEM DOCUMENTATION & USER MANUAL\n")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11)
    run_sub.font.bold = True
    run_sub.font.color.rgb = PURPLE_PRIMARY

    run_title = title_p.add_run("Comprehensive System Documentation & Clinical Guide")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = PURPLE_DARK

    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(20)
    run_meta = meta_p.add_run(
        "Project: Web-Based AI-Assisted Telehealth & Pre-Consultation Safety Platform (MediAI)\n"
        "Lead Engineer: Daniel Isaac Echo (CS/DE/2022/1300) • Department of Computer Science, Caritas University\n"
        "Supervisor & HOD: Mr. Ejike Ugwu • Project Coordinator: Mr. Abundance Akpan • September 2026"
    )
    run_meta.font.name = "Arial"
    run_meta.font.size = Pt(9.5)
    run_meta.font.color.rgb = SLATE_MUTED

    # ----------------------------------------------------
    # CALLOUT BOX: WELCOME & EXECUTIVE OVERVIEW
    # ----------------------------------------------------
    callout_table = doc.add_table(rows=1, cols=1)
    callout_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    callout_cell = callout_table.rows[0].cells[0]
    callout_cell.width = Inches(6.5)
    set_cell_background(callout_cell, HEX_PURPLE_LIGHT)
    set_cell_margins(callout_cell, top=140, bottom=140, left=180, right=180)

    cp = callout_cell.paragraphs[0]
    cp.paragraph_format.space_after = Pt(4)
    run_c_title = cp.add_run("PLATFORM OVERVIEW & CORE CLINICAL MISSION\n")
    run_c_title.font.name = "Arial"
    run_c_title.font.size = Pt(11)
    run_c_title.font.bold = True
    run_c_title.font.color.rgb = PURPLE_DARK

    run_c_body = cp.add_run(
        "MediAI is a unified digital health ecosystem designed to modernize patient-clinician interactions. "
        "The fundamental objective of the platform has been carefully aligned with medical safety standards: "
        "rather than attempting to act as an autonomous drug prescriber (which creates extreme medical malpractice risks such as "
        "prescribing ulcer-eroding NSAIDs to patients with internal stomach bleeding), MediAI operates as a Pre-Consultation "
        "Clinical Triage & Patient Safety Bridge. The system guides patients on immediate safe supportive care ('What to Do and What NOT to Do'), "
        "identifies critical medication contraindications (strictly avoiding NSAIDs like Ibuprofen/Diclofenac/Aspirin), provides ergonomic posture relief, "
        "and seamlessly bridges patients into WhatsApp-style encrypted clinical chats and WebRTC video appointments with verified doctors."
    )
    run_c_body.font.name = "Arial"
    run_c_body.font.size = Pt(9.5)
    run_c_body.font.color.rgb = SLATE_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    def add_section_header(title):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(title)
        r.font.name = "Arial"
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = PURPLE_PRIMARY

    def add_sub_header(title):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(title)
        r.font.name = "Arial"
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = PURPLE_DARK

    def add_body_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.color.rgb = SLATE_DARK
        return p

    def add_bullet_p(title, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r_title = p.add_run(title + ": ")
        r_title.font.name = "Arial"
        r_title.font.size = Pt(10)
        r_title.font.bold = True
        r_title.font.color.rgb = PURPLE_DARK
        r_text = p.add_run(text)
        r_text.font.name = "Arial"
        r_text.font.size = Pt(10)
        r_text.font.color.rgb = SLATE_DARK

    # ----------------------------------------------------
    # SECTION 1: SYSTEM OVERVIEW & WORKSPACE PORTALS
    # ----------------------------------------------------
    add_section_header("1. System Architecture & Role-Based Workspaces")
    add_body_p(
        "MediAI is structured into three specialized user workspaces, engineered with role-based access control (RBAC):"
    )
    add_bullet_p("Patient Portal", 
                 "Provides dynamic symptom selection, interactive clarifying intake questions, multi-condition differential diagnosis, "
                 "critical medication warnings, 1-click specialist appointment booking, encrypted WhatsApp-style doctor messenger, and electronic health record storage.")
    
    add_bullet_p("Doctor Practitioner Workspace", 
                 "Enables medical clinicians to manage consultation schedules, host WebRTC video appointments, submit medical licenses and degree certificates for audit, "
                 "review patient pre-consultation triage memos, and author evidence-based video demonstrations and wellness guides.")
    
    add_bullet_p("Administrator Governance Dashboard", 
                 "Provides the Medical Board audit portal to inspect uploaded doctor certificates, verify licensing numbers, approve or revoke practice credentials, and oversee system security.")

    # ----------------------------------------------------
    # SECTION 2: THE CLINICAL SAFETY PARADIGM & CASE STUDY
    # ----------------------------------------------------
    add_section_header("2. Clinical Safety Triage Engine & Real-Life Case Study")
    add_body_p(
        "Traditional automated healthcare concepts often made the dangerous assumption that algorithms should prescribe drugs directly to patients. "
        "MediAI corrects this flawed approach. The clinical engine functions as a Pre-Consultation Clinical Triage & Patient Safety Bridge."
    )

    add_sub_header("2.1 The Clinical Case Study: Back Pain Locking After Sitting")
    add_body_p(
        "Consider an actual real-world scenario presented to the platform: A female patient experienced severe mid-back pain that internally resembled a previous ulcer episode. "
        "Whenever she sat down for a prolonged period and attempted to stand up, her spine locked up in severe pain, leaving her unable to stand upright for several minutes."
    )
    add_bullet_p("Diagnostic Differential 1 (Referred Gastric Ulcer Pain)", 
                 "When a peptic ulcer or severe gastritis occurs on the posterior wall of the stomach, pain refers through visceral nerves directly to the thoracic back. "
                 "Sitting compresses the stomach; standing stretches the abdominal wall, triggering acute nerve flares that lock the posture.")
    
    add_bullet_p("Diagnostic Differential 2 (Psoas Muscle Spasm)", 
                 "The psoas major originates at the lower spine and travels through the abdomen to the hip. In acute spasm, it causes deep visceral pain that mimics internal organ disease. "
                 "Sitting shortens the muscle; standing forces an immediate stretch against contracted fibers, causing locking pain.")

    add_sub_header("2.2 Critical What to Do and What NOT to Do Rules")
    add_bullet_p("STRICT CONTRAINDICATION (DO NOT TAKE NSAIDs)", 
                 "The patient is strictly warned NEVER to take Ibuprofen, Diclofenac, Feldene, or Aspirin. In suspected ulcers, these drugs inhibit protective prostaglandins, "
                 "causing severe lining erosion and triggering acute internal hemorrhaging.")
    
    add_bullet_p("SAFE OTC ANALGESIC", 
                 "Plain Paracetamol (Acetaminophen) is highlighted as the only safe temporary analgesic prior to seeing a doctor.")
    
    add_bullet_p("ERGONOMIC SUPPORT & CLINIC ESCALATION", 
                 "The patient is advised to sip warm water, avoid peppery foods, and when rising from a chair, shift forward to the chair edge and stand gradually. "
                 "Because the pain is persistent, immediate doctor consultation is recommended.")

    # ----------------------------------------------------
    # SECTION 3: WHATSAPP-GRADE SECURE CLINICAL MESSAGING
    # ----------------------------------------------------
    add_section_header("3. WhatsApp-Grade Secure Clinical Messaging Engine")
    add_body_p(
        "Effective clinical management requires rapid, intuitive doctor-patient communication. MediAI features a secure messaging portal inspired by WhatsApp:"
    )
    add_bullet_p("256-Bit HIPAA End-to-End Encryption Banner", 
                 "Displays a persistent cryptographic privacy notice guaranteeing confidential doctor-patient consultations.")
    add_bullet_p("Double Blue Checkmarks (Read Receipts)", 
                 "Provides real-time delivery and read confirmations using dual blue checkmarks.")
    add_bullet_p("1-Click AI Triage Memo Sharing", 
                 "Patients can send their latest AI assessment directly to the doctor in the chat thread, ensuring the physician has the complete pre-consultation summary immediately.")
    add_bullet_p("Direct Video Call Launch", 
                 "The chat header incorporates a quick 'Video Call' button that seamlessly opens the WebRTC consultation room.")

    # ----------------------------------------------------
    # SECTION 4: FLO-INSPIRED HEALTH COMMUNITY & DOCTOR VIDEOS
    # ----------------------------------------------------
    add_section_header("4. Health Community, Video Demonstrations & Creator Studio")
    add_body_p(
        "MediAI features an evidence-based knowledge hub where verified doctors post clinical guides and video demonstrations:"
    )
    add_bullet_p("Direct Video File Uploads", 
                 "Doctors can upload MP4/WebM video files directly from their phone or computer, demonstrating blood pressure monitoring, physical therapy, or nutrition protocols.")
    add_bullet_p("Clean Semantic Typography (No Hashtags)", 
                 "The article reader strips raw markdown symbols (#, ##, ###), rendering clean, professional headings, bulleted lists, and clinical pearl cards.")
    add_bullet_p("Creator Monetization Studio", 
                 "Verified specialists earn engagement royalties ($0.02/read view, $0.05/like) visible exclusively in their private studio dashboard.")

    # ----------------------------------------------------
    # SECTION 5: DOCTOR CREDENTIAL VERIFICATION PORTAL
    # ----------------------------------------------------
    add_section_header("5. Doctor Credential Verification & Audit System")
    add_body_p(
        "To guarantee patient safety, all medical practitioners must undergo administrative verification:"
    )
    add_bullet_p("Doctor Upload Portal", 
                 "Doctors upload their Medical Practicing Licence, MBBS/MD Degree Certificate, and National ID in their dashboard.")
    add_bullet_p("Admin Document Review", 
                 "Administrators can inspect uploaded documents, preview certificates, verify licence numbers, and grant official practice approval with one click.")

    # ----------------------------------------------------
    # SECTION 6: LIVE DEMONSTRATION GUIDE
    # ----------------------------------------------------
    add_section_header("6. Demonstration Script for Evaluators")
    add_body_p(
        "Follow this sequence to demonstrate MediAI's end-to-end capabilities:"
    )
    add_bullet_p("Step 1 (AI Clinical Triage)", "Navigate to AI Symptom Assessment. Input back pain that locks up when standing from a chair. Observe the differential diagnosis, strict NSAID contraindication, safe Paracetamol recommendation, and 1-click doctor booking.")
    add_bullet_p("Step 2 (WhatsApp Consultation Chat)", "Open the Messenger. Show the 256-bit encryption badge, double checkmarks, and click 'Share AI Triage' to transmit the clinical memo to the doctor.")
    add_bullet_p("Step 3 (Video Education Hub)", "Open Community Feed. View evidence-based articles without raw hashtag symbols, and play uploaded doctor video demonstrations.")
    add_bullet_p("Step 4 (Doctor Verification & Admin Audit)", "Demonstrate the doctor credential upload portal in Doctor Dashboard, then switch to Admin Dashboard to inspect documents and approve credentials.")

    doc.save("MediAI_Comprehensive_System_Documentation.docx")
    print("Saved: MediAI_Comprehensive_System_Documentation.docx")

if __name__ == "__main__":
    create_comprehensive_documentation()
